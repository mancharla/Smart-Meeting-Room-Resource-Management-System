from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter


def create_booking_excel(bookings):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Booking Report"

    headers = [
        "Booking ID",
        "User ID",
        "Room ID",
        "Title",
        "Start Time",
        "End Time",
        "Status",
    ]

    for column, header in enumerate(headers, start=1):
        cell = worksheet.cell(
            row=1,
            column=column,
            value=header,
        )
        cell.font = Font(bold=True)

    for row, booking in enumerate(bookings, start=2):
        worksheet.cell(row=row, column=1, value=booking.id)
        worksheet.cell(row=row, column=2, value=booking.user_id)
        worksheet.cell(row=row, column=3, value=booking.room_id)
        worksheet.cell(row=row, column=4, value=booking.title)
        worksheet.cell(row=row, column=5, value=booking.start_time)
        worksheet.cell(row=row, column=6, value=booking.end_time)
        worksheet.cell(row=row, column=7, value=booking.status)

    # Auto-size columns
    for column_cells in worksheet.columns:
        max_length = 0
        column_letter = get_column_letter(
            column_cells[0].column
        )

        for cell in column_cells:
            if cell.value is not None:
                max_length = max(
                    max_length,
                    len(str(cell.value)),
                )

        worksheet.column_dimensions[
            column_letter
        ].width = min(max_length + 2, 40)

    output = BytesIO()

    workbook.save(output)

    output.seek(0)

    return output